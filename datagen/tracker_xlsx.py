"""The spreadsheet the last accountant left behind.

Every finance team has one. Merged cells, colour as meaning, a header row that
is not the first row, dates typed as text, blank rows where someone hit enter,
and a Notes column holding the only record of several agreements — some still
true, some stale for a year, none dated.

It is the hardest input in the corpus and the most realistic. It is also the
one that must never become a rule: a note saying "abatement thru June per Dana"
is a *claim*, at the bottom of the trust ranking, to be confirmed against a
document before anything is believed. Treating it as a rule is how a system
inherits someone else's mistakes and then compounds them.
"""

from dataclasses import dataclass
from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

TITLE = "Lease tracker - DO NOT DELETE"
HEADER_ROW = 4  # not row 1: there is a title and a blank row above it

GREEN = PatternFill("solid", fgColor="C6EFCE")   # "checked, fine"
AMBER = PatternFill("solid", fgColor="FFEB9C")   # "needs looking at"
RED = PatternFill("solid", fgColor="FFC7CE")     # "problem"

COLOUR_MEANING = {
    "C6EFCE": "reviewed and accepted",
    "FFEB9C": "flagged, not resolved",
    "FFC7CE": "disputed",
}

COLUMNS = ["Clinic", "Landlord", "Suite", "Rent", "CAM est", "Lease exp",
           "Last reviewed", "Notes"]


@dataclass(frozen=True)
class TrackerNote:
    """One row's note, as a low-trust claim rather than a term."""
    clinic_id: str
    note: str
    colour: Optional[str]
    colour_meaning: Optional[str]
    last_reviewed: Optional[str]
    row: int

    @property
    def is_claim(self) -> bool:
        """Always. The tracker never carries authority on its own."""
        return True


# Notes the previous accountant left. Some are still true, some are a year
# stale, and none of them is dated in the cell itself.
_NOTES = {
    "STORE-01": ("Landlord agreed 3% cap verbally - get in writing", AMBER),
    "STORE-02": ("Suite 118 was given back in 2024!! still being billed?", RED),
    "STORE-03": ("escalation date is APRIL not Jan - they keep getting this wrong", AMBER),
    "STORE-04": ("checked against lease 2025, all good", GREEN),
    "STORE-05": ("abatement thru June per Dana - confirm", AMBER),
}


def build_tracker(company, leases) -> bytes:
    """The tracker as an .xlsx file, mess and all."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Leases"

    sheet["A1"] = TITLE
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.merge_cells("A1:H1")
    sheet["A1"].alignment = Alignment(horizontal="left")
    sheet["A2"] = "last updated by M. Okafor - see me before changing anything"
    sheet["A2"].font = Font(italic=True, size=9)
    sheet.merge_cells("A2:H2")

    for column, heading in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=heading)
        cell.font = Font(bold=True)

    row = HEADER_ROW + 1
    for index, terms in enumerate(leases):
        clinic = company.clinic(terms.clinic_id)
        landlord = company.landlord(terms.landlord_id)
        note, fill = _NOTES.get(terms.clinic_id, ("", None))

        sheet.cell(row=row, column=1, value=clinic.clinic_id)
        sheet.cell(row=row, column=2, value=landlord.name)
        sheet.cell(row=row, column=3, value=clinic.suite)
        sheet.cell(row=row, column=4, value=terms.base_rent_monthly)
        sheet.cell(row=row, column=5, value=terms.cam_estimate_monthly)
        # Dates typed as text, in whatever format that day suggested.
        sheet.cell(row=row, column=6, value=terms.expiration.strftime("%m/%d/%y"))
        sheet.cell(row=row, column=7,
                   value="Mar-2025" if index % 2 else "11/14/24")
        sheet.cell(row=row, column=8, value=note)

        if fill is not None:
            for column in range(1, len(COLUMNS) + 1):
                sheet.cell(row=row, column=column).fill = fill

        row += 1
        if index == 2:
            row += 1  # somebody hit enter

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class UnreadableWorkbook(ValueError):
    """The bytes are not a workbook this can read."""


def read_tracker(data: bytes) -> List[TrackerNote]:
    """Pull the notes out as claims.

    The header is not the first row, there are blank rows in the middle, and
    the colour of a row carries meaning nobody wrote down. All of that is
    handled here so no caller has to know about it.
    """
    # The tracker arrives from a shared drive and may be anything. A raw
    # BadZipFile or KeyError halfway through ingestion says nothing useful
    # about which file was bad.
    try:
        book = load_workbook(BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:
        raise UnreadableWorkbook("not a readable .xlsx workbook: %s" % exc) from exc

    sheet = book.active
    if sheet is None:
        raise UnreadableWorkbook("the workbook has no sheets")

    header = _find_header(sheet)
    if header is None:
        return []

    columns = {
        str(sheet.cell(row=header, column=c).value or "").strip().lower(): c
        for c in range(1, sheet.max_column + 1)
    }
    clinic_col = columns.get("clinic")
    notes_col = columns.get("notes")
    reviewed_col = columns.get("last reviewed")
    if clinic_col is None or notes_col is None:
        return []

    out: List[TrackerNote] = []
    for row in range(header + 1, sheet.max_row + 1):
        clinic = sheet.cell(row=row, column=clinic_col).value
        if not clinic:
            continue  # a blank row where somebody hit enter
        note = sheet.cell(row=row, column=notes_col).value
        if not str(note or "").strip():
            continue

        colour = _fill_of(sheet.cell(row=row, column=clinic_col))
        out.append(TrackerNote(
            clinic_id=str(clinic).strip(),
            note=str(note).strip(),
            colour=colour,
            colour_meaning=COLOUR_MEANING.get(colour or ""),
            last_reviewed=(str(sheet.cell(row=row, column=reviewed_col).value).strip()
                           if reviewed_col else None),
            row=row,
        ))
    return out


def _find_header(sheet) -> Optional[int]:
    for row in range(1, min(sheet.max_row, 20) + 1):
        values = {str(sheet.cell(row=row, column=c).value or "").strip().lower()
                  for c in range(1, sheet.max_column + 1)}
        if {"clinic", "notes"} <= values:
            return row
    return None


def _fill_of(cell) -> Optional[str]:
    fill = cell.fill
    if fill is None or fill.fgColor is None:
        return None
    rgb = fill.fgColor.rgb
    if not isinstance(rgb, str) or len(rgb) < 6:
        return None
    colour = rgb[-6:].upper()
    return colour if colour != "000000" else None
